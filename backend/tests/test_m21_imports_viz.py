"""Unit tests for M21 reverse import + viz modules."""
from __future__ import annotations

from pathlib import Path

import pytest


@pytest.fixture(autouse=True)
def _isolated_data_dir(monkeypatch, tmp_path):
    from app import config as _cfg
    monkeypatch.setattr(_cfg, "data_dir", lambda: tmp_path)
    _cfg._CACHE = None
    # ensure /projects/<pid> exists for the importers
    (tmp_path / "projects" / "test_pid").mkdir(parents=True, exist_ok=True)
    yield


def test_protocol_heuristic_extract():
    from app.ingestion.protocol_importer import _heuristic_extract
    txt = ("PROTOCOL NUMBER: NCT-99999\n"
           "Phase II\n"
           "Indication: chronic kidney disease\n"
           "Sample size: 240\n"
           "Primary endpoint: eGFR slope\n")
    m = _heuristic_extract(txt)
    assert m.study_id == "NCT-99999"
    assert "2" in m.phase
    assert "chronic" in m.indication.lower()
    assert m.sample_size == 240
    assert any("egfr" in e.lower() for e in m.endpoints)


def test_sap_importer_no_outline(tmp_path):
    """SAP import should not blow up when project has no outline yet."""
    from docx import Document
    p = tmp_path / "sap.docx"
    doc = Document()
    doc.add_heading("Statistical Methods", level=1)
    doc.add_paragraph("Continuous vars summarised with n, mean, SD.")
    doc.save(str(p))
    from app.ingestion.sap_importer import import_sap
    result = import_sap("test_pid", p, source_filename="sap.docx")
    assert result.n_sections_total >= 1
    assert len(result.analysis_sections) >= 1


def test_define_importer_minimal_xml(tmp_path):
    p = tmp_path / "d.xml"
    p.write_text(
        """<?xml version="1.0" encoding="UTF-8"?>
<ODM xmlns="http://www.cdisc.org/ns/odm/v1.3">
  <Study OID="s1"><MetaDataVersion OID="m1" Name="ADaM">
    <ItemGroupDef OID="IG.ADSL" Name="ADSL">
      <ItemRef ItemOID="IT.SUBJID"/>
    </ItemGroupDef>
    <ItemDef OID="IT.SUBJID" Name="SUBJID" DataType="text" Role="Identifier"/>
  </MetaDataVersion></Study>
</ODM>
""",
        encoding="utf-8",
    )
    from app.ingestion.define_importer import import_define
    res = import_define("test_pid", p)
    assert res.n_datasets >= 1
    assert res.n_variables >= 1


def test_km_with_risk_synthetic(tmp_path):
    import pandas as pd
    import random
    random.seed(1)
    rows = []
    for i in range(40):
        arm = "A" if i % 2 == 0 else "B"
        rows.append({"AVAL": random.uniform(0.5, 14),
                      "CNSR": 0 if random.random() > 0.5 else 1,
                      "TRT01P": arm})
    p = tmp_path / "km.parquet"
    pd.DataFrame(rows).to_parquet(p)
    from app.analysis.advanced.km_with_risk import km_with_risk_table
    block = km_with_risk_table(str(p), "AVAL", "CNSR", "TRT01P",
                                 [0, 3, 6, 9, 12])
    rj = block.result_json
    assert "risk_table" in rj
    assert len(rj["risk_table"]["rows"]) >= 1
    # at-risk decreasing
    for row in rj["risk_table"]["rows"]:
        nars = row["n_at_risk"]
        for a, b in zip(nars, nars[1:]):
            assert a >= b


def test_bland_altman_basic(tmp_path):
    import pandas as pd, random
    random.seed(2)
    rows = []
    for i in range(60):
        m1 = random.uniform(50, 150)
        m2 = m1 + random.uniform(-10, 10)
        rows.append({"M1": m1, "M2": m2})
    p = tmp_path / "ba.parquet"
    pd.DataFrame(rows).to_parquet(p)
    from app.analysis.advanced.bland_altman import bland_altman
    block = bland_altman(str(p), "M1", "M2")
    rj = block.result_json
    assert {"bias", "upper_LoA", "lower_LoA", "n_pairs"}.issubset(rj.keys())
    assert rj["upper_LoA"] >= rj["lower_LoA"]


def test_heatmap_basic(tmp_path):
    import pandas as pd, random
    random.seed(3)
    rows = []
    for g in ("A", "B", "C"):
        for s in range(5):
            rows.append({"G": g, "S": f"s{s}", "V": random.random()})
    p = tmp_path / "hm.parquet"
    pd.DataFrame(rows).to_parquet(p)
    from app.analysis.advanced.heatmap import heatmap
    block = heatmap(str(p), "G", "S", "V")
    rj = block.result_json
    assert len(rj["row_labels"]) == 3
    assert len(rj["col_labels"]) == 5


def test_dashboard_persist(tmp_path):
    (tmp_path / "projects" / "pX").mkdir(parents=True, exist_ok=True)
    from app.analysis.advanced.dashboard import (
        build_dashboard, list_dashboards, delete_dashboard,
    )
    d = build_dashboard("pX", "Demo", [
        {"stat_id": "abc", "position": {"x": 0, "y": 0, "w": 6, "h": 4}},
        {"stat_id": "def", "position": {"x": 6, "y": 0, "w": 6, "h": 4}},
    ])
    assert d.id
    lst = list_dashboards("pX")
    assert any(x.id == d.id for x in lst)
    delete_dashboard("pX", d.id)
    assert all(x.id != d.id for x in list_dashboards("pX"))
