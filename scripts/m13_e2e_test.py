"""End-to-end test for V2-D / M13 — i18n prompts + advanced DOCX + CSR import.

Covers:
  1. Create an EN-language project (language=en).
  2. Verify writer prompt builder picks the en template (offline check).
  3. PATCH /api/projects/<pid>/export/template_config — set Times New Roman,
     body size 12, header text, watermark.
  4. POST /api/projects/<pid>/export/docx in mock pipeline — verify the
     generated .docx zip contains the new font + the header / watermark text.
  5. Build a synthetic 3-H1 + 5-H2 docx, POST /api/projects/<pid>/import/csr
     → ImportResult, then POST commit → outline + drafts persisted.
  6. GET /api/projects/<pid>/state_summary reflects outline/report counts.

Pipeline runs with CSR_WRITER_MOCK=1 so no LLM is exercised end-to-end.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

import requests

ROOT = Path(__file__).resolve().parent.parent
BACKEND = ROOT / "backend"
BASE = "http://127.0.0.1:8781"


def info(msg: str) -> None:
    print(f"[E2E-M13] {msg}", flush=True)


def err(msg: str) -> None:
    print(f"[FAIL] {msg}", file=sys.stderr, flush=True)


def start_server(env: dict) -> subprocess.Popen:
    info("starting uvicorn on :8781")
    args = [
        sys.executable, "-m", "uvicorn",
        "app.server.main:app", "--port", "8781",
        "--host", "127.0.0.1", "--log-level", "warning",
    ]
    proc = subprocess.Popen(
        args, cwd=BACKEND, env=env,
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
    )
    deadline = time.time() + 30
    while time.time() < deadline:
        try:
            r = requests.get(f"{BASE}/api/health", timeout=2)
            if r.status_code == 200:
                info("server up")
                return proc
        except requests.RequestException:
            time.sleep(0.4)
    proc.terminate()
    raise RuntimeError("server did not come up")


def stop_server(proc: subprocess.Popen) -> None:
    try:
        proc.terminate()
        proc.wait(timeout=10)
    except Exception:
        try:
            proc.kill()
        except Exception:
            pass


def cleanup_project(pid: str) -> None:
    try:
        sys.path.insert(0, str(BACKEND))
        from app.config import data_dir
        shutil.rmtree(data_dir() / "projects" / pid, ignore_errors=True)
        pj = data_dir() / "projects.json"
        if pj.exists():
            items = json.loads(pj.read_text(encoding="utf-8") or "[]")
            items = [p for p in items if p.get("id") != pid]
            pj.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass


def build_synthetic_csr(path: Path) -> None:
    """Author a small .docx with 3 H1 + 5 H2 + paragraphs + 1 table."""
    from docx import Document
    doc = Document()
    h1_titles = ["Introduction", "Methods", "Results"]
    h2_under_methods = ["Study Design", "Population", "Statistical Analysis"]
    h2_under_results = ["Demographics", "Efficacy"]

    doc.add_heading(h1_titles[0], level=1)
    doc.add_paragraph("This is the introduction paragraph for the synthetic CSR.")
    doc.add_paragraph("Background and rationale go here.")

    doc.add_heading(h1_titles[1], level=1)
    doc.add_paragraph("Methodology overview.")
    for sub in h2_under_methods:
        doc.add_heading(sub, level=2)
        doc.add_paragraph(f"Body content for {sub} subsection.")

    doc.add_heading(h1_titles[2], level=1)
    for sub in h2_under_results:
        doc.add_heading(sub, level=2)
        doc.add_paragraph(f"Findings for {sub}.")
    # one table at the end of Results
    tbl = doc.add_table(rows=3, cols=2)
    tbl.cell(0, 0).text = "Arm"
    tbl.cell(0, 1).text = "N"
    tbl.cell(1, 0).text = "A"
    tbl.cell(1, 1).text = "100"
    tbl.cell(2, 0).text = "B"
    tbl.cell(2, 1).text = "98"
    doc.save(str(path))


def docx_has_text(docx_path: Path, needles: list[str]) -> dict[str, bool]:
    """Return {needle: present?} by scanning all document.xml + header/footer xml."""
    found = {n: False for n in needles}
    with zipfile.ZipFile(docx_path, "r") as zf:
        for name in zf.namelist():
            if not name.endswith(".xml"):
                continue
            try:
                blob = zf.read(name).decode("utf-8", errors="ignore")
            except Exception:
                continue
            for n in needles:
                if n in blob:
                    found[n] = True
    return found


def main() -> int:
    workdir = Path(tempfile.mkdtemp(prefix="autocsr_m13_e2e_"))
    env = os.environ.copy()
    env["CSR_WRITER_MOCK"] = "1"
    env["CSR_EDITOR_MOCK"] = "1"
    env.setdefault("PYTHONIOENCODING", "utf-8")

    overlay = BACKEND / "app" / "config" / "settings.yaml"
    saved = overlay.read_bytes() if overlay.exists() else None
    overlay.write_text(
        "llm:\n  api_key: ''\n"
        "server:\n  host: 127.0.0.1\n  port: 8781\n",
        encoding="utf-8",
    )

    proc = None
    pid: str | None = None
    try:
        # --- offline prompt loader checks first (no server needed) ---------
        sys.path.insert(0, str(BACKEND))
        from app.i18n.loader import load_prompt, available_prompts
        from app.i18n.terminology import default_zh_en_map, swap_terms
        en_writer = load_prompt("writer", "en")
        assert "Clinical Study Report" in en_writer, "en writer template missing"
        zh_writer = load_prompt("writer", "zh")
        assert "临床研究报告" in zh_writer, "zh writer template missing"
        assert {"writer", "harmonizer", "proposer", "analyst", "editor"}.issubset(
            set(available_prompts("zh")),
        )
        assert {"writer", "harmonizer", "proposer", "analyst", "editor"}.issubset(
            set(available_prompts("en")),
        )
        # bilingual swap
        m = default_zh_en_map()
        swapped = swap_terms("安慰剂组共 50 例不良事件", m)
        assert "placebo" in swapped and "adverse event" in swapped, \
            f"swap_terms failed: {swapped}"
        info("offline prompt + terminology checks passed")

        # --- start server --------------------------------------------------
        proc = start_server(env)

        # --- create English-language project -------------------------------
        r = requests.post(
            f"{BASE}/api/projects",
            json={"name": "e2e_m13_en", "principle_id": "ich_e3",
                  "language": "en"}, timeout=10,
        )
        r.raise_for_status()
        pid = r.json()["id"]
        info(f"created EN project: {pid}")
        # touch GET to populate last_opened_at + verify language echoed
        proj = requests.get(f"{BASE}/api/projects/{pid}", timeout=5).json()
        assert proj["language"] == "en", f"expected language=en, got {proj['language']}"

        # --- template_config: patch font + header + watermark --------------
        r = requests.patch(
            f"{BASE}/api/projects/{pid}/export/template_config",
            json={
                "fonts": {"heading": "Arial", "body": "Times New Roman", "code": "Consolas"},
                "sizes": {"h1": 18, "h2": 16, "h3": 14, "body": 12},
                "colors": {"heading": "#000000", "body": "#000000", "link": "#0000FF"},
                "margins": {"top": 2.54, "bottom": 2.54, "left": 3.17, "right": 3.17},
                "line_spacing": 1.5,
                "toc_depth": 3,
                "header_text": "ICH E3 CSR",
                "footer_text": "Confidential",
                "watermark": "CONFIDENTIAL",
                "show_ai_provenance": False,
            }, timeout=10,
        )
        r.raise_for_status()
        cfg = r.json()
        assert cfg["fonts"]["body"] == "Times New Roman"
        assert cfg["header_text"] == "ICH E3 CSR"
        assert cfg["watermark"] == "CONFIDENTIAL"
        info("template_config patched")

        # --- Need outline + at least one draft to export. Reuse reverse-import path
        # to seed the project quickly. Build a synthetic docx, import, commit. -
        synth = workdir / "synthetic.docx"
        build_synthetic_csr(synth)
        info(f"built synthetic CSR docx: {synth.name}")
        with synth.open("rb") as f:
            r = requests.post(
                f"{BASE}/api/projects/{pid}/import/csr",
                files={"file": (synth.name, f,
                                 "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                timeout=30,
            )
        r.raise_for_status()
        imp = r.json()
        info(f"import: import_id={imp['import_id']} headings={imp['n_headings']} "
             f"paragraphs={imp['n_paragraphs']} tables={imp['n_tables']} "
             f"confidence={imp['confidence']:.2f}")
        # We authored 3 H1 + 5 H2 = 8 headings
        assert imp["n_headings"] >= 8, f"expected ≥8 headings, got {imp['n_headings']}"
        assert imp["n_tables"] >= 1, f"expected ≥1 table, got {imp['n_tables']}"

        r = requests.post(f"{BASE}/api/projects/{pid}/import/csr/commit",
                          json={"import_id": imp["import_id"], "principle_id": "ich_e3"},
                          timeout=20)
        r.raise_for_status()
        commit_res = r.json()
        info(f"commit: n_outline_nodes={commit_res['n_outline_nodes']} "
             f"n_drafts={commit_res['n_drafts']}")
        assert commit_res["n_outline_nodes"] >= 8
        assert commit_res["n_drafts"] >= 5

        # --- state_summary now shows outline + drafts ----------------------
        r = requests.get(f"{BASE}/api/projects/{pid}/state_summary", timeout=10)
        r.raise_for_status()
        s = r.json()
        info(f"state_summary after import: outline={s['steps']['outline']} "
             f"report={s['steps']['report']}")
        assert s["steps"]["outline"]["done"] >= 8
        assert s["steps"]["report"]["done"] >= 5

        # --- build DOCX with the template_config ---------------------------
        r = requests.post(
            f"{BASE}/api/projects/{pid}/export/docx",
            json={
                "include_compliance_note": True,
                "include_toc": True,
                "include_appendix_cleansing": False,
                "include_appendix_analysis": False,
            }, timeout=120,
        )
        r.raise_for_status()
        out = r.json()
        info(f"docx built: {out['filename']} sections={out['n_sections']} "
             f"size={out['size_bytes']}")
        assert out["ok"]

        # download + inspect
        r = requests.get(
            f"{BASE}/api/projects/{pid}/export/docx/{out['filename']}", timeout=30,
        )
        r.raise_for_status()
        docx_path = workdir / out["filename"]
        docx_path.write_bytes(r.content)
        hits = docx_has_text(docx_path, [
            "Times New Roman", "ICH E3 CSR", "CONFIDENTIAL",
        ])
        info(f"docx text presence: {hits}")
        # Header/footer text + watermark stub should all be present in XML
        assert hits["Times New Roman"], "expected font 'Times New Roman' in docx XML"
        assert hits["ICH E3 CSR"], "expected header text 'ICH E3 CSR' in docx XML"
        assert hits["CONFIDENTIAL"], "expected watermark 'CONFIDENTIAL' in docx XML"

        # --- preset listing endpoint --------------------------------------
        r = requests.get(f"{BASE}/api/export/presets", timeout=5)
        r.raise_for_status()
        presets = r.json()
        assert {p["id"] for p in presets} == {"standard", "pharma", "academic", "regulatory"}
        info(f"presets endpoint: {[p['id'] for p in presets]}")

        print("\n[E2E-M13] OK  all checks passed", flush=True)
        return 0
    except AssertionError as e:
        err(f"assertion failed: {e}")
        import traceback; traceback.print_exc()
        return 2
    except Exception as e:
        err(f"exception: {e}")
        import traceback; traceback.print_exc()
        return 3
    finally:
        if proc is not None:
            stop_server(proc)
        if pid:
            cleanup_project(pid)
        try:
            if saved is None:
                if overlay.exists():
                    overlay.unlink()
            else:
                overlay.write_bytes(saved)
        except Exception:
            pass
        try:
            shutil.rmtree(workdir, ignore_errors=True)
        except Exception:
            pass


if __name__ == "__main__":
    sys.exit(main())
